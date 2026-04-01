import os
import json
import time
from datetime import datetime, timedelta

import requests
from dotenv import load_dotenv
from google import genai
from pydantic import BaseModel, Field
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

load_dotenv()

# ── Config ───────────────────────────────────────────────────────────────────

BRIGHT_DATA_DATASET_ID = "gd_lyy3tktm25m4avu764"
GEMINI_MODEL = "gemini-3.1-flash-lite-preview"


# ── Pydantic Model ───────────────────────────────────────────────────────────

class PostSummary(BaseModel):
    summary: str = Field(description="A concise 1-2 sentence summary of the LinkedIn post.")


# ── Bright Data: Fetch LinkedIn Posts ────────────────────────────────────────

def fetch_linkedin_posts(linkedin_url: str, months_back: int) -> list[dict]:
    api_key = os.getenv("BRIGHT_DATA_API_KEY")
    if not api_key:
        print("[!] BRIGHT_DATA_API_KEY not set in .env")
        return []

    end_date = datetime.now().strftime("%Y-%m-%d")
    start_date = (datetime.now() - timedelta(days=months_back * 30)).strftime("%Y-%m-%d")

    url = (
        f"https://api.brightdata.com/datasets/v3/scrape"
        f"?dataset_id={BRIGHT_DATA_DATASET_ID}"
        f"&notify=false"
        f"&include_errors=true"
        f"&type=discover_new"
        f"&discover_by=company_url"
    )
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    payload = json.dumps({
        "input": [{
            "url": linkedin_url,
            "start_date": start_date,
            "end_date": end_date,
        }],
    })

    print(f"Fetching LinkedIn posts ({start_date} to {end_date})...")

    try:
        resp = requests.post(url, headers=headers, data=payload, timeout=300)
        resp.raise_for_status()
    except requests.RequestException as e:
        print(f"[!] Bright Data request failed: {e}")
        if hasattr(e, "response") and e.response is not None:
            print(f"    Response: {e.response.text[:500]}")
        return []

    # Handle NDJSON or regular JSON
    try:
        result = resp.json()
    except requests.exceptions.JSONDecodeError:
        lines = resp.text.strip().split("\n")
        posts = [json.loads(line) for line in lines if line.strip()]
        print(f"Got {len(posts)} posts.")
        return posts

    if isinstance(result, list):
        print(f"Got {len(result)} posts.")
        return result

    # Async response — poll for results
    snapshot_id = result.get("snapshot_id")
    if not snapshot_id:
        print(f"[!] Unexpected response: {result}")
        return []

    snapshot_url = f"https://api.brightdata.com/datasets/v3/snapshot/{snapshot_id}?format=json"
    print(f"  Waiting for results (snapshot: {snapshot_id})...")

    for attempt in range(60):
        time.sleep(5)
        try:
            poll_resp = requests.get(snapshot_url, headers=headers, timeout=30)
            if poll_resp.status_code == 200:
                try:
                    posts = poll_resp.json()
                except requests.exceptions.JSONDecodeError:
                    lines = poll_resp.text.strip().split("\n")
                    posts = [json.loads(line) for line in lines if line.strip()]
                if isinstance(posts, list):
                    print(f"Got {len(posts)} posts.")
                    return posts
            elif poll_resp.status_code == 202:
                if attempt % 6 == 0:
                    print(f"  Still processing... ({attempt * 5}s)")
            else:
                print(f"[!] Poll status {poll_resp.status_code}: {poll_resp.text[:300]}")
                return []
        except requests.RequestException as e:
            print(f"[!] Poll failed: {e}")

    print("[!] Timed out waiting for results.")
    return []


# ── Gemini: Generate Summaries ───────────────────────────────────────────────

def generate_summary(post_text: str) -> str:
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key or not post_text.strip():
        return ""

    client = genai.Client(api_key=api_key)

    resp = client.models.generate_content(
        model=GEMINI_MODEL,
        config={
            "response_mime_type": "application/json",
            "response_schema": PostSummary,
            "system_instruction": (
                "You summarize LinkedIn posts concisely. "
                "Output ONLY valid JSON matching the schema."
            ),
        },
        contents=f"Summarize this LinkedIn post in 1-2 sentences:\n\n{post_text}",
    )

    result = PostSummary.model_validate_json(resp.text)
    return result.summary


# ── Excel Export ─────────────────────────────────────────────────────────────

def save_to_excel(posts: list[dict], linkedin_url: str) -> str:
    os.makedirs("output", exist_ok=True)

    # Extract company name from URL for filename
    company_slug = linkedin_url.rstrip("/").split("/")[-1]
    date_str = datetime.now().strftime("%Y-%m-%d")
    filename = f"output/{company_slug}_posts_{date_str}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "LinkedIn Posts"

    headers = ["Post Date", "Full Content", "Summary", "Post Link"]
    header_font = Font(bold=True, size=12)
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, post in enumerate(posts, 2):
        ws.cell(row=row_idx, column=1, value=post.get("date", ""))
        ws.cell(row=row_idx, column=2, value=post.get("content", ""))
        ws.cell(row=row_idx, column=3, value=post.get("summary", ""))
        ws.cell(row=row_idx, column=4, value=post.get("link", ""))

    # Auto-size columns
    for col_idx, header in enumerate(headers, 1):
        max_len = len(header)
        for row in range(2, len(posts) + 2):
            val = ws.cell(row=row, column=col_idx).value
            if val:
                max_len = max(max_len, min(len(str(val)), 60))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max_len + 4

    wb.save(filename)
    return filename


# ── Main ─────────────────────────────────────────────────────────────────────

def main():
    print("\n📋 LinkedIn Post Exporter")
    print("=" * 40)

    linkedin_url = input("Enter LinkedIn company URL: ").strip()
    months_back = input("How many months back? ").strip()
    months_back = int(months_back) if months_back else 3

    # Step 1: Fetch posts
    raw_posts = fetch_linkedin_posts(linkedin_url, months_back)
    if not raw_posts:
        print("\n⚠️  No posts found.")
        return

    # Step 2: Generate summaries
    print(f"\nGenerating summaries for {len(raw_posts)} posts...")
    posts = []
    for i, raw in enumerate(raw_posts):
        post_text = raw.get("post_text", "")
        print(f"  [{i + 1}/{len(raw_posts)}] Summarizing...")
        summary = generate_summary(post_text)
        # Parse date to clean format
        raw_date = raw.get("date_posted", "")
        try:
            parsed = datetime.fromisoformat(raw_date.replace("Z", "+00:00"))
            clean_date = parsed.strftime("%B %d, %Y")
        except (ValueError, AttributeError):
            clean_date = raw_date or "Unknown"

        posts.append({
            "date": clean_date,
            "content": post_text,
            "summary": summary,
            "link": raw.get("url", ""),
        })

    # Step 3: Export
    filepath = save_to_excel(posts, linkedin_url)
    print(f"\n✅ Exported {len(posts)} posts to {filepath}")


if __name__ == "__main__":
    main()
