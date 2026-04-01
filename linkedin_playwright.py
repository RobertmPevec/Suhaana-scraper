import os
import re
import random
import asyncio
from datetime import datetime

from dotenv import load_dotenv
from playwright.async_api import async_playwright
from google import genai
from pydantic import BaseModel, Field
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

load_dotenv()

GEMINI_MODEL = "gemini-3.1-flash-lite-preview"


# ── Pydantic Model ───────────────────────────────────────────────────────────

class PostSummary(BaseModel):
    summary: str = Field(description="A concise 1-2 sentence summary of the LinkedIn post.")


class BatchSummaries(BaseModel):
    summaries: list[str] = Field(
        description="List of concise 1-2 sentence summaries, one per post in the same order they were provided."
    )


# ── Helper: Parse LinkedIn post age into months ──────────────────────────────

def parse_age_to_months(age_text: str) -> float:
    """Convert LinkedIn age text like '2d', '3w', '1mo', '1yr' to months."""
    age_text = age_text.strip().lower()

    match = re.match(r"(\d+)\s*(d|w|mo|yr|y|h|m)", age_text)
    if not match:
        return 0

    num = int(match.group(1))
    unit = match.group(2)

    if unit in ("h", "m"):
        return 0
    elif unit == "d":
        return num / 30
    elif unit == "w":
        return num / 4
    elif unit == "mo":
        return num
    elif unit in ("yr", "y"):
        return num * 12
    return 0


# ── Scrape LinkedIn Posts with Playwright (Async) ────────────────────────────

async def scrape_linkedin_posts(company_url: str, months_back: int) -> list[dict]:
    """Open LinkedIn, let user log in, then scrape company posts."""

    # Ensure URL ends with /posts/
    posts_url = company_url.rstrip("/")
    if not posts_url.endswith("/posts"):
        posts_url += "/posts/"
    else:
        posts_url += "/"

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=False)
        context = await browser.new_context(
            viewport={"width": 1280, "height": 900},
            user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36",
        )
        page = await context.new_page()

        # Step 1: Go to LinkedIn login
        print("Opening LinkedIn login page...")
        await page.goto("https://www.linkedin.com/login", wait_until="domcontentloaded")
        await asyncio.sleep(random.uniform(1.0, 2.0))

        # Step 2: Auto-login if credentials are in .env, otherwise wait for manual login
        li_email = os.getenv("LINKEDIN_EMAIL", "")
        li_password = os.getenv("LINKEDIN_PASSWORD", "")

        if li_email and li_password:
            print("Logging in automatically...")

            # Type email with random delays
            email_input = page.locator("#username")
            await email_input.click()
            await asyncio.sleep(random.uniform(0.3, 0.7))
            for char in li_email:
                await email_input.type(char, delay=random.randint(50, 150))
            await asyncio.sleep(random.uniform(0.5, 1.0))

            # Type password with random delays
            password_input = page.locator("#password")
            await password_input.click()
            await asyncio.sleep(random.uniform(0.3, 0.7))
            for char in li_password:
                await password_input.type(char, delay=random.randint(40, 130))
            await asyncio.sleep(random.uniform(0.5, 1.0))

            # Click sign in
            await page.locator("button[type='submit']").click()
            print("Credentials submitted. Waiting for redirect...")
        else:
            print("No LINKEDIN_EMAIL/LINKEDIN_PASSWORD in .env. Please log in manually...")

        await page.wait_for_url("**/feed/**", timeout=120000)
        print("Logged in!")

        # Step 3: Navigate to company posts page
        print(f"Navigating to {posts_url}...")
        await page.goto(posts_url, wait_until="domcontentloaded")
        await asyncio.sleep(3)

        # Step 4: Click "Sort by" dropdown and select "Recent"
        print("Sorting by Recent...")
        try:
            sort_button = page.locator("#sort-dropdown-trigger")
            await sort_button.click()
            await asyncio.sleep(1)

            recent_option = page.locator("button[role='option']").filter(has_text="Recent")
            await recent_option.click()
            await asyncio.sleep(2)
            print("Sorted by Recent.")
        except Exception as e:
            print(f"[!] Could not sort by Recent: {e}. Continuing with default sort.")

        # Step 5: Scroll and scrape posts
        print(f"Scraping posts (up to {months_back} months back)...")
        posts = []
        seen_ids = set()
        stop_scraping = False
        scroll_count = 0
        no_new_posts_count = 0

        while not stop_scraping:
            post_elements = await page.locator("div.feed-shared-update-v2").all()

            new_posts_found = 0
            for post_el in post_elements:
                try:
                    urn = await post_el.get_attribute("data-urn") or ""
                    activity_id = ""
                    if "activity:" in urn:
                        activity_id = urn.split("activity:")[-1]

                    if not activity_id:
                        try:
                            link_el = post_el.locator("a[href*='activity']").first
                            href = await link_el.get_attribute("href") or ""
                            match = re.search(r"activity[:-](\d+)", href)
                            if match:
                                activity_id = match.group(1)
                        except:
                            pass

                    if not activity_id or activity_id in seen_ids:
                        continue

                    # Get post age
                    age_text = ""
                    try:
                        age_el = post_el.locator("span.update-components-actor__sub-description span[aria-hidden='true']").first
                        raw_age = (await age_el.inner_text()).strip()
                        age_text = raw_age.split("•")[0].strip()
                    except:
                        pass

                    # Check if we've gone past the time limit
                    if age_text:
                        age_months = parse_age_to_months(age_text)
                        if age_months >= months_back:
                            print(f"  Reached post aged '{age_text}' ({age_months:.1f} months). Stopping.")
                            stop_scraping = True
                            break

                    # Get post text
                    post_text = ""
                    try:
                        text_el = post_el.locator("div.update-components-text span[dir='ltr']").first
                        post_text = (await text_el.inner_text()).strip()
                    except:
                        try:
                            text_el = post_el.locator("div.update-components-text").first
                            post_text = (await text_el.inner_text()).strip()
                        except:
                            pass

                    if not post_text:
                        continue

                    post_url = f"https://www.linkedin.com/feed/update/urn:li:activity:{activity_id}"

                    seen_ids.add(activity_id)
                    new_posts_found += 1
                    posts.append({
                        "text": post_text,
                        "age": age_text,
                        "url": post_url,
                        "activity_id": activity_id,
                    })
                    print(f"  [{len(posts)}] Scraped post ({age_text}): {post_text[:80]}...")

                except Exception as e:
                    continue

            if new_posts_found == 0:
                no_new_posts_count += 1
                if no_new_posts_count >= 10:
                    print("No new posts found after 10 scrolls. Stopping.")
                    break
            else:
                no_new_posts_count = 0

            if stop_scraping:
                break

            # Scroll down realistically
            scroll_count += 1
            scroll_amount = random.randint(800, 1200)
            await page.mouse.wheel(0, scroll_amount)
            wait_time = random.uniform(1.5, 3.0)
            await asyncio.sleep(wait_time)

            if scroll_count % 5 == 0:
                print(f"  Scrolled {scroll_count} times, {len(posts)} posts so far...")

        print(f"\nDone! Scraped {len(posts)} posts.")
        await browser.close()

    return posts


# ── Gemini: Generate Summaries (batch of 5) ──────────────────────────────────

SUMMARY_CHUNK_SIZE = 5


def generate_summaries_batch(post_texts: list[str]) -> list[str]:
    """Generate summaries for up to 5 posts in a single Gemini call."""
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        return [""] * len(post_texts)

    # Build numbered prompt
    numbered_posts = ""
    for i, text in enumerate(post_texts, 1):
        numbered_posts += f"Post {i}:\n{text}\n\n"

    client = genai.Client(api_key=api_key)

    try:
        resp = client.models.generate_content(
            model=GEMINI_MODEL,
            config={
                "response_mime_type": "application/json",
                "response_schema": BatchSummaries,
                "system_instruction": (
                    "You summarize LinkedIn posts concisely. "
                    "Output ONLY valid JSON matching the schema. "
                    "Return exactly one summary per post, in the same order."
                ),
            },
            contents=f"Summarize each of the following {len(post_texts)} LinkedIn posts in 1-2 sentences each:\n\n{numbered_posts}",
        )

        result = BatchSummaries.model_validate_json(resp.text)

        # Ensure we have the right number of summaries
        summaries = result.summaries
        while len(summaries) < len(post_texts):
            summaries.append("")
        return summaries[:len(post_texts)]
    except Exception as e:
        print(f"    [!] Gemini error: {e}. Filling summaries as empty.")
        return [""] * len(post_texts)


# ── Excel Export ─────────────────────────────────────────────────────────────

def save_to_excel(posts: list[dict], company_url: str) -> str:
    os.makedirs("output", exist_ok=True)

    company_slug = company_url.rstrip("/").split("/company/")[-1].split("/")[0]
    date_str = datetime.now().strftime("%Y-%m-%d")
    filename = f"output/{company_slug}_posts_{date_str}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "LinkedIn Posts"

    headers = ["Post Age", "Full Content", "Summary", "Post Link"]
    header_font = Font(bold=True, size=12)
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, post in enumerate(posts, 2):
        ws.cell(row=row_idx, column=1, value=post.get("age", ""))
        ws.cell(row=row_idx, column=2, value=post.get("content", ""))
        ws.cell(row=row_idx, column=3, value=post.get("summary", ""))
        ws.cell(row=row_idx, column=4, value=post.get("link", ""))

    for col_idx, header in enumerate(headers, 1):
        max_len = len(header)
        for row in range(2, len(posts) + 2):
            val = ws.cell(row=row, column=col_idx).value
            if val:
                max_len = max(max_len, min(len(str(val)), 60))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max_len + 4

    wb.save(filename)
    return filename


# ── Standalone Main ──────────────────────────────────────────────────────────

async def _standalone_main():
    print("\n📋 LinkedIn Post Exporter (Playwright)")
    print("=" * 45)

    company_url = input("Enter LinkedIn company URL: ").strip()
    months_back = input("How many months back? ").strip()
    months_back = int(months_back) if months_back else 3
    use_summary = input("Generate Gemini summaries? (yes/no): ").strip().lower() == "yes"

    # Step 1: Scrape posts
    raw_posts = await scrape_linkedin_posts(company_url, months_back)
    if not raw_posts:
        print("\n⚠️  No posts found.")
        return

    # Step 2: Generate summaries if requested
    posts = []
    if use_summary:
        print(f"\nGenerating summaries for {len(raw_posts)} posts ({SUMMARY_CHUNK_SIZE} at a time)...")
        for batch_start in range(0, len(raw_posts), SUMMARY_CHUNK_SIZE):
            batch = raw_posts[batch_start:batch_start + SUMMARY_CHUNK_SIZE]
            batch_num = (batch_start // SUMMARY_CHUNK_SIZE) + 1
            total_batches = (len(raw_posts) + SUMMARY_CHUNK_SIZE - 1) // SUMMARY_CHUNK_SIZE

            print(f"  Batch {batch_num}/{total_batches} ({len(batch)} posts)...")
            texts = [raw["text"] for raw in batch]
            summaries = generate_summaries_batch(texts)

            for raw, summary in zip(batch, summaries):
                posts.append({
                    "age": raw["age"],
                    "content": raw["text"],
                    "summary": summary,
                    "link": raw["url"],
                })
    else:
        print("\nSkipping summaries.")
        for raw in raw_posts:
            posts.append({
                "age": raw["age"],
                "content": raw["text"],
                "summary": "",
                "link": raw["url"],
            })

    # Step 3: Export
    filepath = save_to_excel(posts, company_url)
    print(f"\n✅ Exported {len(posts)} posts to {filepath}")


if __name__ == "__main__":
    asyncio.run(_standalone_main())